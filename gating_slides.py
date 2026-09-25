#!/usr/bin/env python
"""
A FlowJo workspace's layouts (the gating strategy) as PowerPoint slides.

    python3 gating_slides.py WORKSPACE.wsp SAMPLE.fcs

Writes one .pptx per layout in the workspace, one slide each, to outputs/gating/.
Every plot in a layout is redrawn from the FCS events the way FlowJo draws it:
pseudocolor, contour or histogram, on the workspace's own per-channel scales,
with the gates defined on those axes. Axis titles are marker names (CD3, F4/80)
rather than detector channels, at 14 pt.

Each plot is one PowerPoint group, arranged as in the FlowJo layout with room
left for arrows. Inside a group only the events are a picture; the axis titles,
tick labels, plot title, gate names and frequencies are native text, and the
gate outlines are native lines. Gate frequencies are the ones FlowJo saved in
the workspace; flowjo_gating.py re-applies the gating and reports how closely
the redrawn populations match FlowJo's counts.
"""

from __future__ import annotations

import re
import sys
import xml.etree.ElementTree as ET
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from matplotlib.colors import LinearSegmentedColormap
from pptx import Presentation
from pptx.util import Inches
from scipy.ndimage import gaussian_filter, gaussian_filter1d
from scipy.stats import rankdata

import flowjo_gating as fg
import make_slides as ms
import slide_figures as sf

OUTDIR = Path("outputs/gating")
PLOT = 2.25                          # plot area, inches square
MARGINS = (0.92, 0.28, 0.42, 0.72)   # left, right, top, bottom, inches
AXIS_PT, GATE_PT, TITLE_PT = 14, 11, 14
GATE_LINE = "_gate"

# FlowJo's pseudocolor ramp: sparse events blue, dense ones red.
PSEUDO = LinearSegmentedColormap.from_list(
    "flowjo", ["#0000ff", "#0096ff", "#00dc82", "#b4e600", "#ffc800", "#ff6e00", "#ff0000"])

# What the layouts are called in the files, keyed by FlowJo layout name.
FILE_NAMES = {"T Cells": "gating_strategy_T_cells", "Layout-2": "gating_strategy_myeloid"}


# ---------------------------------------------------------------------------
# Names
# ---------------------------------------------------------------------------

def pretty(name: str) -> str:
    """FlowJo population or marker name to display text, signs as superscripts."""
    s = name.replace("F4_80", "F4/80")
    s = re.sub(r"\+-$", "-", s)                     # FlowJo's NOT node, "CD3+-"
    s = re.sub(r"(?<=[A-Za-z0-9])([+-])(?=$|[A-Z ])", r"$^{\1}$", s)
    return s


def axis_title(sample: fg.Sample, channel: str) -> str:
    marker = sample.markers.get(channel, "").strip()
    return pretty(marker) if marker else channel.replace("Comp-", "")


# ---------------------------------------------------------------------------
# Axes
# ---------------------------------------------------------------------------

def _tick_label(v: float) -> str:
    if v == 0:
        return "0"
    k = int(round(np.log10(abs(v))))
    return f"{'−' if v < 0 else ''}$10^{{{k}}}$"


def _label_width(label: str, size: float) -> float:
    plain = ms.text_width(re.sub(r"[${}^]", "", label), size)
    return plain * (0.85 if "^" in label else 1.0)


def biex_ticks(tf: fg.Transform, span_in: float, size: float):
    """Major ticks with labels that do not collide at this size, and minor ticks."""
    lo, hi = tf.data_range
    per_in = tf.length / span_in                    # channels per inch
    priority = [0] + [10 ** k for k in (4, 5, 6, 3, 7, 2)] + [-10 ** k for k in (3, 4, 5, 6)]
    kept: list[tuple[float, float, str]] = []       # (channel, half width in channels, label)
    for v in priority:
        if not lo <= v <= hi:
            continue
        c = float(tf(v))
        if c < 0 or c > tf.length:
            continue
        label = _tick_label(v)
        half = (_label_width(label, size) / 2 + 0.035) * per_in
        if all(abs(c - kc) >= half + kh for kc, kh, _l in kept):
            kept.append((c, half, label))
    kept.sort()
    majors = [c for c, _h, _l in kept]
    labels = [lab for _c, _h, lab in kept]
    # Tick marks for every decade, labelled or not, and minor ticks between.
    candidates = sorted(float(tf(s * m * 10 ** k)) for k in range(1, 8) for m in range(1, 10)
                        for s in (1, -1) if lo <= s * m * 10 ** k <= hi)
    zero = float(tf(0)) if lo <= 0 <= hi else None
    minors: list[float] = []
    for c in candidates:
        if not 0 <= c <= tf.length or any(abs(c - m) < 1.5 for m in majors):
            continue
        # Inside the linear zone around zero the decades pile onto one another.
        if zero is not None and abs(c - zero) < 6:
            continue
        if minors and c - minors[-1] < 1.2:
            continue
        minors.append(c)
    return majors, labels, minors


# ---------------------------------------------------------------------------
# Plots
# ---------------------------------------------------------------------------

def _frame(ax) -> None:
    for side in ("top", "right", "bottom", "left"):
        ax.spines[side].set_color("black")
        ax.spines[side].set_linewidth(0.8)
    ax.tick_params(which="major", length=5, width=0.8, colors="black", labelsize=AXIS_PT,
                   direction="out")
    ax.tick_params(which="minor", length=2.5, width=0.6, colors="black", direction="out")


def _pseudocolor(ax, cx, cy, length):
    h, _xe, _ye = np.histogram2d(cx, cy, bins=length, range=[[0, length], [0, length]])
    h = gaussian_filter(h, sigma=1.6)
    ix = np.clip(cx.astype(int), 0, length - 1)
    iy = np.clip(cy.astype(int), 0, length - 1)
    dens = h[ix, iy]
    shade = rankdata(dens) / len(dens)
    order = np.argsort(dens)
    ax.scatter(cx[order], cy[order], c=shade[order], cmap=PSEUDO, vmin=0, vmax=1, s=0.55,
               marker="s", linewidths=0, rasterized=True, zorder=2)


def _contour(ax, cx, cy, length, level=0.05):
    bins = length // 2
    h, xe, ye = np.histogram2d(cx, cy, bins=bins, range=[[0, length], [0, length]])
    sigma = 2.2 if len(cx) > 1000 else 3.5
    h = gaussian_filter(h, sigma=sigma)
    ix = np.clip((cx / length * bins).astype(int), 0, bins - 1)
    iy = np.clip((cy / length * bins).astype(int), 0, bins - 1)
    dens = h[ix, iy]
    # Probability contours: each line encloses another `level` of the events.
    levels = np.unique(np.quantile(dens, np.arange(level, 1.0, level)))
    levels = levels[levels > 0]
    xc = (xe[:-1] + xe[1:]) / 2
    yc = (ye[:-1] + ye[1:]) / 2
    ax.contour(xc, yc, h.T, levels=levels, colors="black", linewidths=0.55, zorder=2)
    out = dens < levels[0]
    ax.scatter(cx[out], cy[out], s=2.2, c="black", linewidths=0, zorder=2)


def _histogram(ax, cx, length):
    counts, edges = np.histogram(cx, bins=length, range=(0, length))
    smooth = gaussian_filter1d(counts.astype(float), sigma=1.6)
    y = 100 * smooth / smooth.max()
    x = (edges[:-1] + edges[1:]) / 2
    ax.fill_between(x, y, color="#c8c8c8", linewidth=0, zorder=1)
    ax.plot(x, y, color="black", linewidth=0.9, zorder=2)
    return x, y


def _crowding_grid(kind: str, cx, cy, lx: float, ly: float, curve=None, n: int = 64):
    """How full each part of the plot is, 0 to 1, on an n x n grid indexed [x, y]."""
    if kind == "Histogram":
        xs = (np.arange(n) + 0.5) / n * lx
        ys = (np.arange(n) + 0.5) / n * ly
        return (ys[None, :] <= np.interp(xs, *curve)[:, None]).astype(float)
    h, _xe, _ye = np.histogram2d(cx, cy, bins=n, range=[[0, lx], [0, ly]])
    h = gaussian_filter(h, 1.0)
    # Occupied: the region holding the densest 90% of the events, the area a
    # contour plot's outer line would enclose. Text over stray events reads fine.
    ix = np.clip((cx / lx * n).astype(int), 0, n - 1)
    iy = np.clip((cy / ly * n).astype(int), 0, n - 1)
    level = np.quantile(h[ix, iy], 0.10) if len(cx) else np.inf
    return (h >= level).astype(float)


def _crowding(grid, box, lx: float, ly: float) -> float:
    n = grid.shape[0]
    i0, i1 = int(max(0, box[0] / lx * n)), int(min(n, np.ceil(box[2] / lx * n)))
    j0, j1 = int(max(0, box[1] / ly * n)), int(min(n, np.ceil(box[3] / ly * n)))
    return float(grid[i0:i1, j0:j1].mean()) if i1 > i0 and j1 > j0 else 0.0


def _overlap(a, b) -> bool:
    return a[0] < b[2] and a[2] > b[0] and a[1] < b[3] and a[3] > b[1]


def _inside(a, b) -> bool:
    return a[0] >= b[0] and a[2] <= b[2] and a[1] >= b[1] and a[3] <= b[3]


def _candidates(rect, w: float, h: float, lx: float, ly: float):
    """Label centres around and inside a gate, most FlowJo-like first."""
    x0, y0, x1, y1 = rect
    gx, gy = 0.012 * lx, 0.012 * ly
    out = []
    xs = np.linspace(x0 - w / 2, x1 + w / 2, 13)
    xs = sorted(xs, key=lambda x: abs(x - (x0 + x1) / 2))
    for step in (0.5, 1.6):
        out += [(x, y1 + gy + step * h, "out") for x in xs]
        out += [(x, y0 - gy - step * h, "out") for x in xs]
    ys = sorted(np.linspace(y0 + h / 2, y1 - h / 2, 5) if y1 - y0 > h else [(y0 + y1) / 2],
                key=lambda y: -y)
    for x in (x1 + gx + w / 2, x0 - gx - w / 2):
        out += [(x, y, "out") for y in ys]
    for x in (x0 + gx + w / 2, (x0 + x1) / 2, x1 - gx - w / 2):
        out += [(x, y1 - gy - h / 2, "in"), (x, y0 + gy + h / 2, "in")]
    return out


def _gap(box, rect) -> float:
    """Distance between a label box and its gate, 0 when touching or inside."""
    dx = max(rect[0] - box[2], box[0] - rect[2], 0)
    dy = max(rect[1] - box[3], box[1] - rect[3], 0)
    return float(np.hypot(dx, dy))


def place_labels(gates: list[tuple[tuple, tuple]], grid, lx: float, ly: float):
    """Centre for each gate's label: clear of the other labels, off the gate outlines
    and the densest events, and near its own gate. Every placing order is tried and
    the best arrangement kept. gates: [(rect x0 y0 x1 y1, (w, h))], plot units."""
    from itertools import permutations
    rects = [r for r, _wh in gates]
    cands = [_candidates(r, w, h, lx, ly) for r, (w, h) in gates]

    def place(order):
        placed, centres, total = [], {}, 0.0
        for i in order:
            rect, (w, h) = gates[i]
            best = None
            for k, (xc, yc, where) in enumerate(cands[i]):
                # Slide a label that would run off the plot back inside it.
                xc = min(max(xc, 0.01 * lx + w / 2), 0.99 * lx - w / 2)
                yc = min(max(yc, 0.01 * ly + h / 2), 0.99 * ly - h / 2)
                box = (xc - w / 2, yc - h / 2, xc + w / 2, yc + h / 2)
                if where == "in" and not _inside(box, rect):
                    continue
                if any(_overlap(box, b) for b in placed):
                    continue
                score = 4 * _crowding(grid, box, lx, ly) + 0.002 * k
                score += sum(1.5 for r in rects if _overlap(box, r) and not _inside(box, r))
                score += 0.8 * _gap(box, rect) / h
                if best is None or score < best[0]:
                    best = (score, box, xc, yc)
            if best is None:        # nowhere clear: just inside the top of the gate
                xc, yc = (rect[0] + rect[2]) / 2, rect[3] - 0.012 * ly - h / 2
                best = (10.0, (xc - w / 2, yc - h / 2, xc + w / 2, yc + h / 2), xc, yc)
            placed.append(best[1])
            centres[i] = (best[2], best[3])
            total += best[0]
        return total, centres

    orders = permutations(range(len(gates))) if len(gates) <= 5 else [range(len(gates))]
    _total, centres = min((place(o) for o in orders), key=lambda t: t[0])
    return [centres[i] for i in range(len(gates))]


def _label_size(text: str, lx: float, ly: float) -> tuple[float, float]:
    """Label width and height in plot units."""
    lines = sf.markup(text).split("\n")
    w_in = max(ms.text_width(line, GATE_PT) for line in lines) + 0.06
    h_in = len(lines) * GATE_PT * 1.1 / 72 + 0.03
    return w_in * lx / PLOT, h_in * ly / PLOT


def _gate_label(pop: fg.Population, parent: fg.Population) -> str:
    pct = 100 * pop.count / parent.count
    freq = f"{pct:.3g}" if pct >= 1 else f"{pct:.2g}"
    return f"{pretty(pop.name)}\n{freq}"


def draw_plot(sample: fg.Sample, spec: dict, name: str) -> sf.Panel:
    pop = sample.populations[spec["path"]]
    xch, ych, kind = spec["x"], spec["y"], spec["type"]
    left, right, top, bottom = MARGINS
    w, h = left + PLOT + right, top + PLOT + bottom
    fig = plt.figure(figsize=(w, h))
    ax = fig.add_axes([left / w, bottom / h, PLOT / w, PLOT / h])
    tx = sample.transforms[xch]
    L = tx.length
    cx = tx(sample.events[xch][pop.mask])

    hist_curve = None
    if kind == "Histogram":
        hist_curve = _histogram(ax, cx, L)
        ax.set_ylim(0, 105)
        ax.set_yticks([0, 25, 50, 75, 100])
        ax.set_yticklabels(["0", "25", "50", "75", "100"])
        ylabel = "% of Max"
    else:
        ty = sample.transforms[ych]
        cy = ty(sample.events[ych][pop.mask])
        (_pseudocolor if kind == "Pseudocolor" else _contour)(ax, cx, cy, L)
        ax.set_ylim(0, ty.length)
        yt, yl, ym = biex_ticks(ty, PLOT, AXIS_PT)
        ax.set_yticks(yt)
        ax.set_yticklabels(yl)
        ax.set_yticks(ym, minor=True)
        ylabel = axis_title(sample, ych)
    ax.set_xlim(0, L)
    xt, xl, xm = biex_ticks(tx, PLOT, AXIS_PT)
    ax.set_xticks(xt)
    ax.set_xticklabels(xl)
    ax.set_xticks(xm, minor=True)
    _frame(ax)
    ax.set_xlabel(axis_title(sample, xch), fontsize=AXIS_PT, color="black", labelpad=4)
    ax.set_ylabel(ylabel, fontsize=AXIS_PT, color="black", labelpad=4)
    ax.set_title(pretty(pop.name), fontsize=TITLE_PT, fontweight="bold", color="black", pad=7)

    # Gates drawn on these axes, the way FlowJo shows them: positive gates only.
    ly = 105 if kind == "Histogram" else sample.transforms[ych].length
    gates = []
    for child in pop.children:
        g = child.gate
        if g is None or child.negated or g.kind != "rect":
            continue
        if kind == "Histogram":
            if g.dims != [xch]:
                continue
            x0, x1 = (float(tx(v)) if v is not None else e
                      for v, e in zip((g.lo[0], g.hi[0]), (0, L)))
            yr = 100 * g.extra.get("yRatio", 0.75)
            x0, x1 = max(x0, 0), min(x1, L)
            ax.plot([x0, x0, x1, x1], [yr - 5, yr, yr, yr - 5], color="black", linewidth=1.2,
                    label=GATE_LINE, clip_on=False, zorder=4)
            gates.append(((x0, yr - 5, x1, yr), _gate_label(child, pop)))
            continue
        if set(g.dims) != {xch, ych}:
            continue
        span = {d: (lo, hi) for d, lo, hi in zip(g.dims, g.lo, g.hi)}
        ty = sample.transforms[ych]
        x0, x1 = (float(tx(v)) if v is not None else e for v, e in zip(span[xch], (0, L)))
        y0, y1 = (float(ty(v)) if v is not None else e for v, e in zip(span[ych], (0, ty.length)))
        x0, x1 = max(x0, 0), min(x1, L)
        y0, y1 = max(y0, 0), min(y1, ty.length)
        ax.plot([x0, x1, x1, x0, x0], [y0, y0, y1, y1, y0], color="black", linewidth=1.2,
                label=GATE_LINE, clip_on=False, zorder=4)
        gates.append(((x0, y0, x1, y1), _gate_label(child, pop)))

    if gates:
        grid = _crowding_grid(kind, cx, None if kind == "Histogram" else cy, L, ly,
                              curve=hist_curve)
        sized = [(rect, _label_size(text, L, ly)) for rect, text in gates]
        for (rect, text), (w_, h_), (xc, yc) in zip(gates, [wh for _r, wh in sized],
                                                     place_labels(sized, grid, L, ly)):
            # A label that could only go over events gets a light backing to stay legible.
            box = (xc - w_ / 2, yc - h_ / 2, xc + w_ / 2, yc + h_ / 2)
            kind_ = "gate_backed" if _crowding(grid, box, L, ly) > 0.08 else "gate"
            ax.text(xc, yc, text, ha="center", va="center", fontsize=GATE_PT, color="black",
                    gid=kind_, linespacing=1.05, zorder=5)

    base = OUTDIR / re.sub(r"[^A-Za-z0-9]+", "_", f"{spec['layout']}_{spec['index']}_{pop.name}")
    return sf.lift(fig, ax, name, base, line_labels=(GATE_LINE,))


# ---------------------------------------------------------------------------
# Layouts and slides
# ---------------------------------------------------------------------------

def layouts(wsp: Path) -> list[dict]:
    root = ET.parse(wsp).getroot()
    out = []
    for lay in root.iter("Layout"):
        plots = []
        for i, cd in enumerate(lay.find("FigList").findall("ChartData")):
            g = cd.find("Graph")
            axes = {a.get("dimension"): a.get("name") for a in g.findall("Axis")}
            wp = cd.find("WindowPosition")
            dl = cd.find("PopModelList").find("DataLayer")
            plots.append({"layout": lay.get("name"), "index": i, "type": g.get("type"),
                          "x": axes.get("x"), "y": axes.get("y") or None, "path": dl.get("path"),
                          "sample": dl.get("sampleID"),
                          "pos": (float(wp.get("x")), float(wp.get("y")))})
        out.append({"name": lay.get("name"), "plots": plots})
    return out


def arrange(plots: list[dict], cw: float, ch: float) -> list[tuple[float, float]]:
    """Slide positions that keep the FlowJo layout's rows and columns, spaced for arrows."""
    ys = sorted({round(p["pos"][1] / 60) for p in plots})
    rows = [[p for p in plots if round(p["pos"][1] / 60) == y] for y in ys]
    for r in rows:
        r.sort(key=lambda p: p["pos"][0])
    ncol = max(len(r) for r in rows)
    col_gap = min(2.0, (ms.W - 0.6 - ncol * cw) / max(ncol - 1, 1))
    row_gap = min(0.5, (ms.H - 0.3 - len(rows) * ch) / max(len(rows) - 1, 1))
    x0 = (ms.W - ncol * cw - (ncol - 1) * col_gap) / 2
    y0 = (ms.H - len(rows) * ch - (len(rows) - 1) * row_gap) / 2
    pos = {}
    for ri, r in enumerate(rows):
        for ci, p in enumerate(r):
            pos[p["index"]] = (x0 + ci * (cw + col_gap), y0 + ri * (ch + row_gap))
    return [pos[p["index"]] for p in plots]


def build(wsp: Path, fcs: Path) -> list[Path]:
    sample = fg.load(wsp, fcs)
    written = []
    for lay in layouts(wsp):
        plots = [p for p in lay["plots"] if p["sample"] == sample.sample_id]
        if not plots:
            continue
        panels = []
        for p in plots:
            pop = sample.populations[p["path"]]
            axes_txt = axis_title(sample, p["x"]) + (
                f" vs {axis_title(sample, p['y'])}" if p["y"] else "")
            label = sf.markup(f"{pretty(pop.name)}: {axes_txt}").replace("^{", "").replace("}", "")
            panels.append(draw_plot(sample, p, label))
        cw, ch = panels[0].width, panels[0].height
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(ms.W), Inches(ms.H)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.notes_slide.notes_text_frame.text = (
            f"FlowJo layout '{lay['name']}', sample {sample.name}. Gate frequencies are % of the "
            "parent population, as saved in the workspace. Axis titles are marker names; "
            "channels: " + "; ".join(
                f"{axis_title(sample, p['x'])} = {p['x'].replace('Comp-', '')}" for p in plots) + ".")
        for (x, y), panel in zip(arrange(plots, cw, ch), panels):
            ms.add_panel(slide, panel, x, y, 1.0, panel.letter)
        out = OUTDIR / f"{FILE_NAMES.get(lay['name'], re.sub(r'[^A-Za-z0-9]+', '_', lay['name']))}.pptx"
        prs.save(out)
        written.append(out)
    return written


def main(argv=None) -> int:
    argv = argv if argv is not None else sys.argv[1:]
    if len(argv) != 2:
        print(__doc__)
        return 2
    OUTDIR.mkdir(parents=True, exist_ok=True)
    for out in build(Path(argv[0]), Path(argv[1])):
        print(f"Saved: {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
